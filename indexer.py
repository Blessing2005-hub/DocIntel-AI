import os
import json

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

try:
    import pytesseract
    from PIL import Image
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False



BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)


UPLOAD_FOLDER = os.path.join(
    BASE_DIR,
    "uploads"
)


INDEX_FILE = os.path.join(
    BASE_DIR,
    "document_index.json"
)



# -------------------------
# TEXT EXTRACTION
# -------------------------

def extract_text(file_path):

    text = ""

    extension = os.path.splitext(
        file_path
    )[1].lower()



    # PDF

    if extension == ".pdf":

        reader = PdfReader(
            file_path
        )

        for page in reader.pages:

            text += (
                page.extract_text()
                or ""
            )



    # WORD

    elif extension in [".docx", ".doc"]:

        doc = Document(
            file_path
        )

        for paragraph in doc.paragraphs:

            text += (
                paragraph.text
                + "\n"
            )



    # EXCEL

    elif extension in [".xlsx", ".xls"]:

        workbook = load_workbook(
            file_path,
            data_only=True
        )

        for sheet in workbook:

            for row in sheet.iter_rows(
                values_only=True
            ):

                text += " ".join(
                    [
                        str(cell)
                        for cell in row
                        if cell is not None
                    ]
                )

                text += "\n"



    # POWERPOINT

    elif extension in [".pptx", ".ppt"]:

        presentation = Presentation(
            file_path
        )

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(
                    shape,
                    "text"
                ):

                    text += (
                        shape.text
                        + "\n"
                    )



    # TEXT FILE

    elif extension == ".txt":

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            text = file.read()



    # CSV

    elif extension == ".csv":

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            text = file.read()



    # IMAGE OCR

    elif extension in [
        ".png",
        ".jpg",
        ".jpeg",
        ".webp"
    ]:


        if OCR_AVAILABLE:

            image = Image.open(
                file_path
            )

            text = pytesseract.image_to_string(
                image
            )


        else:

            text = (
                "Image uploaded. "
                "OCR is not installed."
            )



    return text



# -------------------------
# BUILD INDEX
# -------------------------

def rebuild_index():


    documents = []



    if not os.path.exists(
        UPLOAD_FOLDER
    ):

        return 0



    supported = [

        ".pdf",
        ".doc",
        ".docx",
        ".xlsx",
        ".xls",
        ".ppt",
        ".pptx",
        ".txt",
        ".csv",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp"

    ]



    for filename in os.listdir(
        UPLOAD_FOLDER
    ):


        extension = os.path.splitext(
            filename
        )[1].lower()



        if extension in supported:


            path = os.path.join(
                UPLOAD_FOLDER,
                filename
            )


            content = extract_text(
                path
            )



            documents.append({

                "filename": filename,

                "content": content

            })



    with open(
        INDEX_FILE,
        "w",
        encoding="utf-8"
    ) as file:


        json.dump(
            documents,
            file,
            indent=4,
            ensure_ascii=False
        )



    return len(documents)